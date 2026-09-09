"""DOCX in standard manuscript format — the deliverable for a human editor.

Deliberately plain: 12pt Times New Roman, double-spaced, one-inch margins,
chapters on fresh pages. This is not a designed book, it is a document built
for Track Changes and margin comments.
"""

from __future__ import annotations

import re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from . import config as c
from .sources import Book, Chapter, _plain_text

# Standard manuscript marks a scene break with a centred hash, not the book's
# printed ornament.
MS_SCENE_BREAK = "#"

ITALIC_RE = re.compile(r"(<em>.*?</em>|<strong>.*?</strong>)", re.S)
TAG_RE = re.compile(r"<[^>]+>")


def _add_rich(par, html: str) -> None:
    """Add text to a paragraph, preserving <em>/<strong> as real runs."""
    for piece in ITALIC_RE.split(html):
        if not piece:
            continue
        text = _plain_text(piece)
        if not text:
            continue
        run = par.add_run(text)
        run.italic = piece.startswith("<em>")
        run.bold = piece.startswith("<strong>")


def _chapter(doc, ch: Chapter) -> None:
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.paragraph_format.space_before = Pt(72)
    head.paragraph_format.space_after = Pt(24)
    head.add_run(ch.heading_plain).bold = True

    for i, scene in enumerate(ch.scenes):
        if i:
            br = doc.add_paragraph(MS_SCENE_BREAK)
            br.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for block in re.findall(r"<p[^>]*>(.*?)</p>", scene, re.S):
            par = doc.add_paragraph()
            par.paragraph_format.first_line_indent = Inches(0.5)
            _add_rich(par, block)


def build_docx(book: Book) -> None:
    c.BUILD_DIR.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)

    for section in doc.sections:
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Inches(1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(c.TITLE.upper()).bold = True
    for line in (c.SUBTITLE, c.AUTHOR_DISPLAY, f"{book.word_count:,} words"):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for i, ch in enumerate(book.chapters):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _chapter(doc, ch)

    doc.save(c.OUT_DOCX)
